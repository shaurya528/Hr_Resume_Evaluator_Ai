import os
from pathlib import Path
from dotenv import load_dotenv
from groq import Groq
from pydantic import BaseModel
import json
from pypdf import PdfReader
from docx import Document
import time

load_dotenv()

My_Api_Key=os.getenv("GROQ_API_KEY")

if not My_Api_Key:
    raise ValueError("api key is missing")

Model=os.getenv("Model")
if not Model:
    raise ValueError("model is missing ")
Client=Groq(api_key=My_Api_Key)
class Jobd(BaseModel):
    role: str
    required_skills: list[str]
    preferred_skills: list[str]
    minimum_experience: float | None
    education_requirements: list[str]
    responsibilities: list[str]

Job_Schema=Jobd.model_json_schema()
User_Prompt="""We are looking for an enthusiastic AI Intern who wants to gain practical experience in Artificial Intelligence, Machine Learning, and AI-based application development. The intern will work on real-world projects and collaborate with the technical team.

Key Responsibilities

Assist in developing and testing AI/ML-based applications.
Work with Python for data processing and AI development.
Explore and integrate AI APIs and open-source AI models.
Work on chatbots, automation, and AI-powered applications.
Assist with data collection, preprocessing, and analysis.
Implement basic Machine Learning and Generative AI concepts.
Work with APIs, databases, and AI tools.
Debug, test, and improve AI applications.
Collaborate with team members and contribute to project development.
Document project work and technical processes.
Required Skills

Basic knowledge of Python.
Understanding of AI/ML fundamentals.
Familiarity with APIs and JSON.
Basic knowledge of Git/GitHub.
Problem-solving and analytical skills.
Willingness to learn new AI technologies and tools.
Good to Have

Knowledge of Generative AI, LLMs, RAG, LangChain, or n8n.
Familiarity with OpenAI or open-source models such as Ollama.
Basic knowledge of SQL, FastAPI, or Django.
Experience with GitHub projects.
Eligibility

Students and freshers are eligible to apply.
Candidates pursuing or recently completed B.Tech, B.E., BCA, MCA, M.Tech, or related programs can apply.
Candidates with a basic understanding of Python and AI/ML are preferred.
Prior internship or professional experience is not mandatory.
Candidates with a strong interest in learning and building AI projects are encouraged to apply.
What You Will Gain

Hands-on experience with real-world AI projects.
Practical exposure to AI and automation workflows.
Experience working with development tools and APIs.
Mentorship and guidance from the technical team.
Internship certificate and project experience.
Opportunity to build a strong AI-focused portfolio."""

System_Prompt=f"""
Role: act as a hr for a IT Company
Task: your task is fetch the job description and give the response accoding to the {Job_Schema}
Constraints:Fetch all detail from the Job description not invent anything,Do NOT return the schema itself.Do NOT return fields like "properties", "title" or "type".
Fill the schema with actual information extracted from the job description.Fill the schema with actual information extracted from the job description.
If minimum experience is not mentioned, return null.
If information for a list is missing, return an empty list.
Output: give me response in json and accoding to {Job_Schema}
"""
response=Client.chat.completions.create(model=Model,messages=[{
    "role":"system",
    "content":System_Prompt
},{"role":"user",
    "content":User_Prompt}
    ],response_format={
        "type":'json_object'
    })


answer=response.choices[0].message.content 
Raw_Json=json.loads(answer)
job_data=Jobd(**Raw_Json)

class Matchresult(BaseModel):
    score:float
    details:dict

class Experience(BaseModel):
    company: str | None = None
    role: str | None = None
    duration: str | None = None
    description: str | None = None
    skills_used: list[str] = []

class Resume(BaseModel):
    name: str | None = None
    email: str | None = None
    phone: str | None = None

    total_experience_years: float | None = None

    skills: list[str] = []
    experiences: list[Experience] = []
    education: list[str] = []
    projects: list[str] = []
    certifications: list[str] = []

resume_schema = Resume.model_json_schema()


def Parse_resume(text_file):
    System_prompt=f"""   You are an expert resume parser.

    Extract information from the resume based on its meaning,
    not only based on exact section headings.

    Different resumes may use different headings.

    For example:
    - Experience
    - Professional Experience
    - Work History
    - Employment
    - Internships

    These may all contain relevant experience.

    Skills may also appear in the skills section, work experience,
    internships or projects.

    Return ONLY valid JSON matching this schema:
    {resume_schema}

    Important rules:

    1. Do not invent information.
    2. If a value is not available, return null.
    3. If a list has no information, return an empty list.
    4. Include internships inside experiences.
    5. Extract skills mentioned across the entire resume.
    
    """
    response=Client.chat.completions.create(model=Model,messages=[{
    "role":"system",
    "content":System_prompt
    },{"role":"user",
    "content":f""" parse this resume \n{text_file}"""}
    ],response_format={
        "type":'json_object'
    })

    answer=response.choices[0].message.content
    raw=json.loads(answer)
    resume_data=Resume(**raw)
    return resume_data


def final_score(job,resume):

    match_schema = Matchresult.model_json_schema()
    prompt = f"""
    You are an HR recruiter.

    Compare the candidate's resume with the job description.

    JOB DESCRIPTION:
    {job.model_dump_json(indent=2)}

    CANDIDATE RESUME:
    {resume.model_dump_json(indent=2)}
    Return JSON matching this schema:

    {match_schema}

    Give me:
     score for the resume and detail of the resume

    Keep the response concise and easy to read.
    """
    message={
        "role": "user",
        "content" : prompt
    }
    messages=[message]
    response_format={
        "type": "json_object"
    }
    response = Client.chat.completions.create(model=Model, messages=messages, response_format=response_format)
    data = json.loads(response.choices[0].message.content)
    return Matchresult(**data)

def read_pdf(file_path):
    reader = PdfReader(file_path)
    text = ""
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + "\n"
    return text

def read_docx(file_path):
    document = Document(file_path)
    text = ""
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + "\n"
    
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text += cell.text + "\n"
    return text


def read_resume(file_path):
    if file_path.suffix.lower() == ".pdf":
        return read_pdf(file_path)
    elif file_path.suffix.lower() == ".docx":
        return read_docx(file_path)
    else:
        return None


Resume_folder = Path("resumes")
all_results = []

if not Resume_folder.is_dir():
    raise FileNotFoundError("resumes folder not found")

for file_path in Resume_folder.iterdir():

    if file_path.suffix.lower() not in [".pdf", ".docx"]:
        continue

    print("\nProcessing:", file_path.name)

    try:
        resume_text = read_resume(file_path)

        if not resume_text or not resume_text.strip():
            print("No readable text found")
            continue

        parsed_resume = Parse_resume(resume_text)

        time.sleep(5)

        result = final_score(job_data, parsed_resume)

        time.sleep(5)

        print("Candidate:", parsed_resume.name)
        print("Score:", result.score)
        

        all_results.append({
            "score":result.score,
            "details": result.details
        })

    except Exception as e:
        print(f"Error processing {file_path.name}: {e}")

all_results.sort(
    key=lambda candidate: candidate["score"],
    reverse=True
)

print("\nFinal Ranking")


